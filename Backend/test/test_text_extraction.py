"""
Tests for src/text_extraction.py.

Fixtures are real documents built in memory with the same libraries used to
read them, so the tests exercise actual file parsing rather than a mock. No
network and no API key are needed, which keeps the suite runnable in CI.
"""

import io

import pytest

from src.text_extraction import (
    SUPPORTED_EXTENSIONS,
    UnsupportedFileType,
    extract_segments,
    extract_text,
    is_extractable,
)


# -- fixtures ---------------------------------------------------------------

def make_pdf(pages: list[str]) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_font("Helvetica", size=12)
    for page in pages:
        pdf.add_page()
        pdf.multi_cell(0, 8, page)
    return bytes(pdf.output())


def make_docx(blocks: list[tuple[str, str]]) -> bytes:
    """blocks: list of (kind, content); kind is 'heading', 'para' or 'table'."""
    from docx import Document

    document = Document()
    for kind, content in blocks:
        if kind == "heading":
            document.add_heading(content, level=1)
        elif kind == "para":
            document.add_paragraph(content)
        elif kind == "table":
            rows = [line.split(",") for line in content.split(";")]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pptx(slides: list[tuple[str, str]]) -> bytes:
    """slides: list of (body, notes)."""
    from pptx import Presentation

    presentation = Presentation()
    blank = presentation.slide_layouts[5]
    for body, notes in slides:
        slide = presentation.slides.add_slide(blank)
        slide.shapes.title.text = body
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def make_xlsx(sheets: dict[str, list[list]]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(title=name)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# -- dispatch ---------------------------------------------------------------

@pytest.mark.parametrize("filename", [
    "a.pdf", "a.docx", "a.pptx", "a.xlsx", "a.csv", "a.txt", "a.md", "A.PDF",
])
def test_is_extractable_accepts_supported_formats(filename):
    assert is_extractable(filename)


@pytest.mark.parametrize("filename", ["a.png", "a.exe", "a.zip", "a.doc", "noextension", ""])
def test_is_extractable_rejects_others(filename):
    assert not is_extractable(filename)


def test_unsupported_extension_raises_with_a_useful_message():
    with pytest.raises(UnsupportedFileType) as excinfo:
        extract_segments("photo.png", b"\x89PNG")
    assert ".png" in str(excinfo.value)
    # The message must name the alternatives, otherwise the caller can only
    # report failure without telling the user what would have worked.
    assert ".docx" in str(excinfo.value)


def test_every_supported_extension_has_an_extractor():
    for extension in SUPPORTED_EXTENSIONS:
        assert is_extractable(f"file{extension}")


# -- plain text -------------------------------------------------------------

def test_plain_text_roundtrips():
    segments = extract_segments("notes.txt", "Hello world.\n\nSecond paragraph.".encode())
    assert len(segments) == 1
    assert "Second paragraph." in segments[0].text
    assert segments[0].location is None


def test_empty_file_yields_no_segments():
    assert extract_segments("empty.txt", b"") == []
    assert extract_segments("blank.txt", b"   \n\n  ") == []


def test_cp1252_umlauts_survive():
    """A lossy decode would turn 'Betriebsanleitung für Räume' into text with
    replacement characters exactly where the searchable German words are."""
    data = "Betriebsanleitung für Räume".encode("cp1252")
    text = extract_text("anleitung.txt", data)
    assert "für" in text and "Räume" in text
    assert "\ufffd" not in text


# -- csv --------------------------------------------------------------------

def test_csv_rows_stay_on_one_line():
    """Chunkers split on line and paragraph boundaries, so a record broken
    across two lines can end up split across two chunks and match nothing."""
    data = b"room,co2,temp\nExperience Hub,412,21.5\nServer Room,388,18.0\n"
    text = extract_text("readings.csv", data)
    lines = text.splitlines()
    assert lines[0] == "room | co2 | temp"
    assert "Experience Hub | 412 | 21.5" in lines


def test_csv_semicolon_delimiter_is_detected():
    """European exports use ';' because ',' is the decimal separator; reading
    those with a ',' delimiter produces one unusable cell per row."""
    data = "Raum;CO2;Temperatur\nExperience Hub;412;21,5\n".encode()
    lines = extract_text("messwerte.csv", data).splitlines()
    assert lines[0] == "Raum | CO2 | Temperatur"
    assert "Experience Hub | 412 | 21,5" in lines


# -- pdf --------------------------------------------------------------------

def test_pdf_segments_are_labelled_by_page():
    data = make_pdf(["First page about thermostats.", "Second page about valves."])
    segments = extract_segments("manual.pdf", data)
    assert [s.location for s in segments] == ["page 1", "page 2"]
    assert "thermostats" in segments[0].text
    assert "valves" in segments[1].text


def test_pdf_pages_without_text_are_skipped_not_emitted_empty():
    """An image-only page yields nothing. Emitting it as an empty segment
    would cost an embedding and add a chunk that can never be relevant."""
    data = make_pdf(["Only this page has text.", "   "])
    segments = extract_segments("mixed.pdf", data)
    assert len(segments) == 1
    assert segments[0].location == "page 1"


# -- docx -------------------------------------------------------------------

def test_docx_segments_split_at_headings():
    data = make_docx([
        ("heading", "Installation"),
        ("para", "Mount the unit on a flat wall."),
        ("heading", "Maintenance"),
        ("para", "Replace the filter every six months."),
    ])
    segments = extract_segments("guide.docx", data)
    assert len(segments) == 2
    assert segments[0].location == "section: Installation"
    assert segments[1].location == "section: Maintenance"
    assert "filter" in segments[1].text


def test_docx_table_stays_with_its_section():
    """python-docx exposes paragraphs and tables as separate sequences, so a
    naive reader appends every table at the end of the document and severs it
    from the text that introduces it."""
    data = make_docx([
        ("heading", "Intervals"),
        ("para", "The service intervals are as follows."),
        ("table", "Part,Interval;Filter,6 months;Belt,24 months"),
        ("heading", "Contact"),
        ("para", "Call the operator."),
    ])
    segments = extract_segments("intervals.docx", data)
    intervals = next(s for s in segments if s.location == "section: Intervals")
    assert "Filter | 6 months" in intervals.text
    contact = next(s for s in segments if s.location == "section: Contact")
    assert "Filter" not in contact.text


def test_docx_without_headings_yields_one_unlocated_segment():
    data = make_docx([("para", "Just a note."), ("para", "And another.")])
    segments = extract_segments("note.docx", data)
    assert len(segments) == 1
    assert segments[0].location is None


# -- pptx -------------------------------------------------------------------

def test_pptx_segments_are_labelled_by_slide_and_include_notes():
    """Slide bodies are keyword fragments; the explanation is in the notes.
    Indexing only the body produces chunks that match nothing."""
    data = make_pptx([
        ("Air handling overview", "The unit runs at 60% during office hours."),
        ("Filter classes", ""),
    ])
    segments = extract_segments("deck.pptx", data)
    assert [s.location for s in segments] == ["slide 1", "slide 2"]
    assert "60% during office hours" in segments[0].text


# -- xlsx -------------------------------------------------------------------

def test_xlsx_segments_are_labelled_by_sheet():
    data = make_xlsx({
        "Budget": [["Item", "Cost"], ["Server", 1200]],
        "Staff": [["Name", "Role"], ["Ayse", "Engineer"]],
    })
    segments = extract_segments("plan.xlsx", data)
    assert [s.location for s in segments] == ["sheet Budget", "sheet Staff"]
    assert "Server | 1200" in segments[0].text


def test_xlsx_empty_sheets_are_skipped():
    data = make_xlsx({"Data": [["a", 1]], "Blank": []})
    segments = extract_segments("sparse.xlsx", data)
    assert [s.location for s in segments] == ["sheet Data"]


def test_xlsx_trailing_empty_cells_do_not_pad_the_line():
    data = make_xlsx({"S": [["a", "b", None, None], ["c", None, None, None]]})
    lines = extract_text("padded.xlsx", data).splitlines()
    assert lines == ["a | b", "c"]


# -- extract_text -----------------------------------------------------------

def test_extract_text_joins_all_segments():
    data = make_pdf(["Page one text.", "Page two text."])
    text = extract_text("doc.pdf", data)
    assert "Page one text." in text and "Page two text." in text


def test_extract_text_matches_segment_contents():
    data = make_xlsx({"A": [["x", 1]], "B": [["y", 2]]})
    segments = extract_segments("two.xlsx", data)
    text = extract_text("two.xlsx", data)
    for segment in segments:
        assert segment.text in text