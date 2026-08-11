"""
Tests for src/rag/service.py.

End-to-end through the real pipeline -- extraction, chunking, storage,
retrieval -- with only the embedding call faked. Documents are built in
memory, so nothing here needs a server or an API key.
"""

import io
import warnings

import pytest
from qdrant_client import AsyncQdrantClient

from src.rag.embedding import EmbeddingError
from src.rag.service import (
    RagConfig,
    RagService,
    format_context,
)
from src.rag.store import DocumentStore, StoredChunk

VECTOR_SIZE = 8


@pytest.fixture
def anyio_backend():
    """Pin to asyncio; see test_rag_store.py."""
    return "asyncio"


class FakeEmbedder:
    def __init__(self, size: int = VECTOR_SIZE, fail: bool = False):
        self.size = size
        self.fail = fail

    def _vector(self, text):
        vector = [0.0] * self.size
        for character in text.lower():
            if character.isalnum():
                vector[ord(character) % self.size] += 1.0
        norm = sum(v * v for v in vector) ** 0.5 or 1.0
        return [v / norm for v in vector]

    async def embed_documents(self, texts):
        if self.fail:
            raise EmbeddingError("provider unavailable")
        return [self._vector(text) for text in texts]

    async def embed_query(self, text):
        return self._vector(text)

    async def dimensions(self):
        return self.size


@pytest.fixture
async def service():
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message=".*Payload indexes have no effect.*")
        client = AsyncQdrantClient(":memory:")
        store = DocumentStore(client, vector_size=VECTOR_SIZE)
        yield RagService(store, FakeEmbedder(), config=RagConfig(min_index_chars=200))
        await client.close()


def long_text(topic: str, characters: int = 1200) -> bytes:
    """A document long enough to pass the size gate."""
    sentence = f"This paragraph explains {topic} in the context of building operations. "
    body = ""
    while len(body) < characters:
        body += sentence
    return body.encode()


def make_docx(blocks: list[tuple[str, str]]) -> bytes:
    from docx import Document

    document = Document()
    for kind, content in blocks:
        if kind == "heading":
            document.add_heading(content, level=1)
        else:
            document.add_paragraph(content)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# -- the size gate ----------------------------------------------------------

@pytest.mark.anyio
async def test_small_documents_are_refused_with_a_reason(service, anyio_backend):
    """Refusing is the better answer, not a limitation: retrieval would hand
    the model fragments of a document it could have read whole."""
    result = await service.index_document("s1", "f1", "note.txt", b"A short note.")
    assert not result.indexed
    assert "small enough to read directly" in result.reason
    assert result.chunks == 0


@pytest.mark.anyio
async def test_large_documents_are_indexed(service, anyio_backend):
    result = await service.index_document("s1", "f1", "manual.txt", long_text("thermostats"))
    assert result.indexed
    assert result.chunks > 0


@pytest.mark.anyio
async def test_unsupported_formats_are_refused_with_alternatives(service, anyio_backend):
    result = await service.index_document("s1", "f1", "photo.png", b"\x89PNG")
    assert not result.indexed
    assert ".docx" in result.reason


@pytest.mark.anyio
async def test_a_file_with_no_text_is_reported_distinctly(service, anyio_backend):
    """A scanned PDF and an unsupported format need different remedies, so
    they must not produce the same message."""
    result = await service.index_document("s1", "f1", "scan.txt", b"    \n\n   ")
    assert not result.indexed
    assert "OCR" in result.reason


# -- indexing and searching -------------------------------------------------

@pytest.mark.anyio
async def test_indexed_documents_become_searchable(service, anyio_backend):
    await service.index_document("s1", "f1", "manual.txt", long_text("thermostat resets"))
    result = await service.search("s1", "thermostat")
    assert result
    assert "thermostat" in result.context.lower()


@pytest.mark.anyio
async def test_search_before_indexing_returns_nothing(service, anyio_backend):
    result = await service.search("s1", "anything")
    assert not result
    assert result.sources == []


@pytest.mark.anyio
async def test_reindexing_does_not_duplicate(service, anyio_backend):
    data = long_text("filters")
    await service.index_document("s1", "f1", "manual.txt", data)
    first = await service.list_documents("s1")
    await service.index_document("s1", "f1", "manual.txt", data)
    second = await service.list_documents("s1")
    assert first[0]["chunks"] == second[0]["chunks"]


@pytest.mark.anyio
async def test_embedding_failure_is_reported_not_raised(service, anyio_backend):
    """A provider outage should tell the user what happened, not surface as a
    traceback in the middle of a chat."""
    service.embedder = FakeEmbedder(fail=True)
    result = await service.index_document("s1", "f1", "manual.txt", long_text("valves"))
    assert not result.indexed
    assert "Could not embed" in result.reason


# -- attribution ------------------------------------------------------------

def test_format_context_numbers_and_attributes_passages():
    """Without a handle, an instruction to cite sources produces either
    nothing or an invented filename."""
    chunks = [
        StoredChunk(text="Hold the button.", file_id="f1", filename="manual.pdf",
                    chunk_index=3, location="page 4", score=0.9),
        StoredChunk(text="Replace the filter.", file_id="f2", filename="guide.docx",
                    chunk_index=1, location="section: Wartung", score=0.8),
    ]
    context, sources = format_context(chunks)

    assert "[1] manual.pdf, page 4" in context
    assert "[2] guide.docx, section: Wartung" in context
    assert [source["number"] for source in sources] == [1, 2]
    assert sources[0]["file_id"] == "f1"


def test_format_context_omits_a_missing_location():
    chunks = [StoredChunk(text="Body.", file_id="f", filename="notes.txt",
                          chunk_index=0, location=None, score=0.5)]
    context, _ = format_context(chunks)
    assert "[1] notes.txt\n" in context
    assert "None" not in context


@pytest.mark.anyio
async def test_search_carries_locations_through_the_whole_pipeline(service, anyio_backend):
    """Extraction records where text came from; the point is that it survives
    chunking and storage and reaches the citation."""
    data = make_docx([
        ("heading", "Maintenance intervals"),
        ("para", "The filter must be replaced every six months. " * 40),
    ])
    await service.index_document("s1", "f1", "guide.docx", data)
    result = await service.search("s1", "filter replacement")

    assert result.sources
    assert any("Maintenance intervals" in (source["location"] or "")
               for source in result.sources)


# -- budget -----------------------------------------------------------------

@pytest.mark.anyio
async def test_context_is_trimmed_to_the_budget(service, anyio_backend):
    """SAGE already spends around 13k tokens on a simple question before any
    document is involved."""
    await service.index_document("s1", "f1", "manual.txt", long_text("filters", 20_000))

    generous = await service.search("s1", "filters", budget_tokens=4000)
    tight = await service.search("s1", "filters", budget_tokens=200)

    assert len(tight.context) < len(generous.context)
    assert tight.sources, "the best passage must survive any budget"


# -- document management ----------------------------------------------------

@pytest.mark.anyio
async def test_deactivated_documents_drop_out_of_search(service, anyio_backend):
    await service.index_document("s1", "f1", "manual.txt", long_text("thermostats"))
    await service.set_active("s1", "f1", False)
    assert not await service.search("s1", "thermostat")


@pytest.mark.anyio
async def test_reactivation_needs_no_reindexing(service, anyio_backend):
    await service.index_document("s1", "f1", "manual.txt", long_text("thermostats"))
    await service.set_active("s1", "f1", False)
    await service.set_active("s1", "f1", True)
    assert await service.search("s1", "thermostat")


@pytest.mark.anyio
async def test_removing_a_document_removes_it_from_search(service, anyio_backend):
    await service.index_document("s1", "f1", "manual.txt", long_text("thermostats"))
    await service.remove_document("s1", "f1")
    assert not await service.search("s1", "thermostat")
    assert await service.list_documents("s1") == []


@pytest.mark.anyio
async def test_clearing_a_session_leaves_others_intact(service, anyio_backend):
    await service.index_document("alice", "f1", "a.txt", long_text("thermostats"))
    await service.index_document("bob", "f2", "b.txt", long_text("thermostats"))

    await service.clear_session("alice")

    assert not await service.search("alice", "thermostat")
    assert await service.search("bob", "thermostat")


@pytest.mark.anyio
async def test_a_removed_document_disappears_from_lexical_search(service, anyio_backend):
    """The lexical index is cached per session and rebuilt only on
    invalidation. Left stale after a removal, keyword search keeps returning
    passages from a document that no longer exists -- verified: without
    invalidation the deleted text is still retrieved.
    """
    await service.index_document("s1", "f1", "a.txt", long_text("thermostats"))
    await service.index_document("s1", "f2", "b.txt", long_text("Fehler 22 filters"))
    await service.search("s1", "Fehler")          # populates the cache

    await service.remove_document("s1", "f2")

    assert "Fehler" not in (await service.search("s1", "Fehler")).context


@pytest.mark.anyio
async def test_a_newly_indexed_document_is_lexically_searchable(service, anyio_backend):
    """The other half: a stale cache makes a new document visible to dense
    search and invisible to keyword search, which reads as retrieval being
    unreliable rather than as a bug."""
    await service.index_document("s1", "f1", "a.txt", long_text("thermostats"))
    await service.search("s1", "thermostat")      # populates the cache

    await service.index_document("s1", "f2", "b.txt", long_text("Fehler 22 filters"))

    result = await service.search("s1", "Fehler")
    assert "Fehler" in result.context


# -- configuration ----------------------------------------------------------

def test_config_reads_the_environment(monkeypatch):
    """Configuration comes from the environment so the evaluation is a sweep
    rather than a series of edits."""
    monkeypatch.setenv("RAG_MIN_INDEX_CHARS", "500")
    monkeypatch.setenv("ENABLE_HYBRID_SEARCH", "false")
    monkeypatch.setenv("ENABLE_RERANKING", "true")

    config = RagConfig.from_env()
    assert config.min_index_chars == 500
    assert config.hybrid is False
    assert config.rerank is True


def test_config_falls_back_on_nonsense(monkeypatch):
    """A typo in a deployment variable should not take the backend down."""
    monkeypatch.setenv("RAG_MIN_INDEX_CHARS", "not-a-number")
    assert RagConfig.from_env().min_index_chars == 20_000


def test_config_defaults_are_the_documented_ones():
    config = RagConfig()
    assert config.hybrid is True
    assert config.rerank is False