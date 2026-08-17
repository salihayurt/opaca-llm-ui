"""
Plain-text extraction from uploaded documents.

SAGE accepts uploads of any type (the backend has no allow-list; see the
skipped `test_append_files_disallowed`), but only PDFs and images ever reach a
model: `file_utils.upload_files` branches on `is_pdf`/`is_image` and silently
skips everything else. This module supplies the missing half -- turning a
DOCX, PPTX, XLSX, CSV or plain-text upload into text that can be indexed for
retrieval or handed to a model directly.

Two design points worth knowing before changing anything here:

1. Extraction returns *located* segments, not one flat string. A retrieved
   passage has to be attributable to a place in the source document ("page 4",
   "slide 12", "sheet Budget"), otherwise a citation can only name the file.
   Formats that genuinely have no such structure return a single segment with
   `location=None` rather than inventing one.

2. Tabular content is emitted one row per line with cells joined by " | ".
   Chunkers split on blank lines, so keeping a row on a single line stops a
   record from being cut in half between two chunks, which would leave both
   halves unmatchable.

Everything here is a pure function over `bytes`: no session, no network, no
LLM. That keeps it fully testable offline, which matters because the CI runs
without an API key.
"""

from __future__ import annotations

import csv
import io
import logging
from dataclasses import dataclass
from pathlib import Path

from .text_repair import build_vocabulary, looks_shattered, repair_spacing

logging.getLogger(__name__)
logger = logging.getLogger(__name__)


class UnsupportedFileType(Exception):
    """Raised when no extractor is registered for a file's extension."""


@dataclass(frozen=True)
class TextSegment:
    """A piece of a document together with where it came from.

    `location` is a short human-readable label ("page 3", "slide 2",
    "sheet Budget") suitable for showing in a citation, or None when the
    format has no meaningful internal structure to point at.
    """

    text: str
    location: str | None = None


# Extensions that carry no formatting and are read as-is.
PLAIN_TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".json", ".log", ".yaml", ".yml"}


def _decode(data: bytes) -> str:
    """Decode bytes to str, tolerating non-UTF-8 uploads.

    German documents exported from older Windows tooling are frequently
    cp1252, where a UTF-8 strict decode fails outright and `errors="replace"`
    turns every umlaut into a replacement character -- which then breaks
    keyword matching on exactly the words a user would search for. Trying
    cp1252 before falling back to lossy decoding keeps those words intact.
    """
    for encoding in ("utf-8", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _row_to_line(cells: list[str]) -> str:
    """Join one table row into a single line, dropping trailing empty cells."""
    cleaned = [(c or "").strip().replace("\n", " ") for c in cells]
    while cleaned and not cleaned[-1]:
        cleaned.pop()
    return " | ".join(cleaned)


def _extract_plain_text(data: bytes) -> list[TextSegment]:
    text = _decode(data).strip()
    return [TextSegment(text=text)] if text else []


def _extract_csv(data: bytes) -> list[TextSegment]:
    """Normalise a CSV into ' | '-joined rows.

    The delimiter is sniffed rather than assumed: European exports commonly
    use ';' because ',' is the decimal separator, and reading those with a
    ',' delimiter yields one giant unusable cell per row.
    """
    text = _decode(data)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel

    lines = [_row_to_line(row) for row in csv.reader(io.StringIO(text), dialect)]
    body = "\n".join(line for line in lines if line)
    return [TextSegment(text=body)] if body.strip() else []


def _extract_pdf(data: bytes) -> list[TextSegment]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))

    if reader.is_encrypted:
        # An empty user password is common for "print-protected" PDFs and
        # decrypts fine; a real password is not something we can resolve here.
        try:
            if not reader.decrypt(""):
                raise UnsupportedFileType("PDF is password protected")
        except (NotImplementedError, ValueError) as e:
            raise UnsupportedFileType(f"PDF cannot be decrypted: {e}") from e

    segments = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception as e:
            # One malformed page must not lose the other 200. Scanned or
            # image-only pages legitimately yield nothing; both are skipped.
            logger.warning("Could not extract page %d: %s", number, e)
            continue
        if text:
            segments.append(TextSegment(text=text, location=f"page {number}"))
    return segments


def _iter_docx_blocks(document):
    """Yield a DOCX's paragraphs and tables in document order.

    python-docx exposes `document.paragraphs` and `document.tables` as two
    separate sequences, so reading them one after the other puts every table
    at the end of the text regardless of where it appeared. For a document
    where a table follows the paragraph explaining it, that separation
    destroys the association a retriever needs. Walking the body XML restores
    the real order.
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            yield Paragraph(child, document)
        elif tag == "tbl":
            yield Table(child, document)


def _extract_docx(data: bytes) -> list[TextSegment]:
    """Extract a DOCX, segmented by heading.

    DOCX has no page concept available without rendering, so headings are used
    as the location label instead -- "Section: Maintenance intervals" is at
    least as useful in a citation as a page number, and often more so.
    """
    from docx import Document
    from docx.table import Table

    document = Document(io.BytesIO(data))

    segments: list[TextSegment] = []
    heading: str | None = None
    buffer: list[str] = []

    def flush():
        body = "\n\n".join(part for part in buffer if part.strip())
        if body.strip():
            location = f"section: {heading}" if heading else None
            segments.append(TextSegment(text=body.strip(), location=location))
        buffer.clear()

    for block in _iter_docx_blocks(document):
        if isinstance(block, Table):
            rows = [_row_to_line([cell.text for cell in row.cells]) for row in block.rows]
            table_text = "\n".join(row for row in rows if row)
            if table_text:
                buffer.append(table_text)
            continue

        text = block.text.strip()
        if not text:
            continue

        style = (block.style.name or "") if block.style is not None else ""
        if style.startswith("Heading") or style == "Title":
            flush()
            heading = text
            buffer.append(text)
        else:
            buffer.append(text)

    flush()
    return segments


def _extract_pptx(data: bytes) -> list[TextSegment]:
    """Extract a PPTX slide by slide, including speaker notes.

    Notes are included because presentation slides are typically keyword
    fragments while the actual explanation lives in the notes -- indexing only
    the slide body tends to produce chunks that match nothing.
    """
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))

    segments = []
    for number, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False):
                rows = [_row_to_line([cell.text for cell in row.cells]) for row in shape.table.rows]
                table_text = "\n".join(row for row in rows if row)
                if table_text:
                    parts.append(table_text)

        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notes:
                parts.append(f"Notes: {notes}")

        body = "\n\n".join(parts).strip()
        if body:
            segments.append(TextSegment(text=body, location=f"slide {number}"))
    return segments


def _extract_xlsx(data: bytes) -> list[TextSegment]:
    """Extract a workbook one segment per sheet.

    `data_only=True` reads the values Excel last calculated rather than the
    formula source: a cell containing `=SUM(B2:B40)` is worth indexing as
    "18400", not as its formula.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)

    segments = []
    try:
        for sheet in workbook.worksheets:
            lines = []
            for row in sheet.iter_rows(values_only=True):
                line = _row_to_line(["" if value is None else str(value) for value in row])
                if line:
                    lines.append(line)
            body = "\n".join(lines).strip()
            if body:
                segments.append(TextSegment(text=body, location=f"sheet {sheet.title}"))
    finally:
        workbook.close()
    return segments


EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".pptx": _extract_pptx,
    ".xlsx": _extract_xlsx,
    ".xlsm": _extract_xlsx,
    ".csv": _extract_csv,
    ".tsv": _extract_csv,
    **{extension: _extract_plain_text for extension in PLAIN_TEXT_EXTENSIONS},
}

SUPPORTED_EXTENSIONS = frozenset(EXTRACTORS)


def is_extractable(filename: str) -> bool:
    """Whether text can be extracted from a file, judged by its name alone."""
    return Path(filename or "").suffix.lower() in SUPPORTED_EXTENSIONS


def _repair_if_shattered(segments: list[TextSegment]) -> list[TextSegment]:
    """Rejoin words that the PDF extractor split on kerning gaps.

    Some publishers position individual characters with separate drawing
    operations, and extractors read the gaps as word boundaries. EUR-Lex does
    this: the GDPR comes out as 'REGUL A TIONS ... Ar ticle 33 ... Notif
    ication of a personal dat a breach'.

    Left alone, such a document contains no token 'article' and no token
    'notification', so lexical search finds neither and BM25 contributes
    nothing to hybrid retrieval. The vocabulary is rebuilt from the document
    itself, so no word list ships with this.

    Applied only when the text looks shattered, so an ordinary document is not
    put through a transformation it does not need.
    """
    combined = "\n".join(segment.text for segment in segments)
    if not looks_shattered(combined):
        return segments

    logger.info("Text appears split by kerning; repairing word spacing")
    vocabulary = build_vocabulary(combined)
    return [
        TextSegment(text=repair_spacing(segment.text, vocabulary),
                    location=segment.location)
        for segment in segments
    ]


def extract_segments(filename: str, data: bytes) -> list[TextSegment]:
    """Extract `data` into located segments, dispatching on the file extension.

    Raises UnsupportedFileType if the extension has no extractor. Returns an
    empty list for a file that is readable but holds no text -- a scanned PDF
    with no text layer, or an empty spreadsheet. Callers must distinguish the
    two: the first is a user error, the second is a document that simply
    cannot be indexed and should be reported as such rather than silently
    producing nothing.
    """
    extension = Path(filename or "").suffix.lower()
    extractor = EXTRACTORS.get(extension)
    if extractor is None:
        raise UnsupportedFileType(
            f"No text extractor for '{extension or filename}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    return _repair_if_shattered(extractor(data))


def extract_text(filename: str, data: bytes, separator: str = "\n\n") -> str:
    """Extract `data` as one plain string, discarding segment locations.

    For callers that only need the content -- passing a small document to a
    model as text, say. Anything that will later cite the source should use
    extract_segments instead and keep the locations.
    """
    return separator.join(segment.text for segment in extract_segments(filename, data))